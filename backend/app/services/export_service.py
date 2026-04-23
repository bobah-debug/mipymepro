import os
import io
from datetime import datetime
from decimal import Decimal
from sqlalchemy.orm import Session
from app.config import settings
from app.models.product import Product
from app.models.sale import Sale, SaleStatus


def export_inventory_excel(db: Session) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = [
        "Código", "Código de Barras", "Nombre", "Categoría", "Proveedor",
        "Precio Compra", "Precio Venta", "Precio Venta c/IVA",
        "Stock Actual", "Stock Mínimo", "Unidad", "Afecto IVA", "Activo"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    products = db.query(Product).order_by(Product.name).all()

    for row_idx, p in enumerate(products, 2):
        low_stock = p.stock_current <= p.stock_minimum
        row_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid") if low_stock else None

        values = [
            p.code, p.barcode, p.name,
            p.category.name if p.category else "",
            p.supplier.name if p.supplier else "",
            float(p.purchase_price), float(p.sale_price), float(p.sale_price_with_iva),
            p.stock_current, p.stock_minimum, p.unit,
            "Sí" if p.afecto_iva else "No",
            "Sí" if p.is_active else "No",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            if row_fill:
                cell.fill = row_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_inventory_csv(db: Session) -> str:
    import csv
    import io

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Código", "Código Barras", "Nombre", "Categoría", "Proveedor",
        "Precio Compra", "Precio Venta", "Precio Venta c/IVA",
        "Stock Actual", "Stock Mínimo", "Unidad", "Afecto IVA"
    ])

    products = db.query(Product).filter(Product.is_active == True).order_by(Product.name).all()
    for p in products:
        writer.writerow([
            p.code, p.barcode or "", p.name,
            p.category.name if p.category else "",
            p.supplier.name if p.supplier else "",
            str(p.purchase_price), str(p.sale_price), str(p.sale_price_with_iva),
            p.stock_current, p.stock_minimum, p.unit,
            "SI" if p.afecto_iva else "NO"
        ])

    return output.getvalue()


def export_sales_excel(db: Session, year: int, month: int) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from sqlalchemy import extract

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Ventas {month:02d}-{year}"

    headers = [
        "ID", "Folio", "Fecha", "Estado", "Cliente RUT", "Cliente Nombre",
        "Subtotal", "IVA", "Total", "Método Pago", "DTE Enviado"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    sales = db.query(Sale).filter(
        extract("year", Sale.created_at) == year,
        extract("month", Sale.created_at) == month,
        Sale.status != SaleStatus.ANULADA
    ).order_by(Sale.created_at).all()

    for row_idx, s in enumerate(sales, 2):
        ws.cell(row=row_idx, column=1, value=s.id)
        ws.cell(row=row_idx, column=2, value=s.folio)
        ws.cell(row=row_idx, column=3, value=s.created_at.strftime("%d/%m/%Y %H:%M"))
        ws.cell(row=row_idx, column=4, value=s.status.value)
        ws.cell(row=row_idx, column=5, value=s.customer_rut or "")
        ws.cell(row=row_idx, column=6, value=s.customer_name or "Consumidor Final")
        ws.cell(row=row_idx, column=7, value=float(s.subtotal))
        ws.cell(row=row_idx, column=8, value=float(s.iva_amount))
        ws.cell(row=row_idx, column=9, value=float(s.total))
        ws.cell(row=row_idx, column=10, value=s.payment_method.name if s.payment_method else "")
        ws.cell(row=row_idx, column=11, value="Sí" if s.dte_enviado else "No")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
